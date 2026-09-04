from __future__ import annotations

import json
import re
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook
from PIL import Image


SOURCE = Path(r"F:\Math\Drive")
WORK = Path(r"C:\Users\Robin\source\repos\Gojo\.codex-tmp\gojo-slides-20260901")
MEDIA = WORK / "extracted_media"

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}


def safe_name(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip("-.")
    return value[:100] or "source"


def xml_text(blob: bytes, namespace: str, tag: str) -> list[str]:
    root = ET.fromstring(blob)
    values = []
    for node in root.findall(f".//{{{namespace}}}{tag}"):
        if node.text:
            values.append(node.text)
    return values


def image_metadata(path: Path) -> dict:
    result = {"path": str(path), "bytes": path.stat().st_size}
    try:
        with Image.open(path) as image:
            result.update({"width": image.width, "height": image.height, "mode": image.mode})
    except Exception as exc:
        result["inspection_error"] = str(exc)
    return result


def extract_package_media(path: Path, prefixes: tuple[str, ...]) -> list[dict]:
    target = MEDIA / safe_name(str(path.relative_to(SOURCE)))
    target.mkdir(parents=True, exist_ok=True)
    found = []
    with zipfile.ZipFile(path) as package:
        for name in package.namelist():
            if not any(name.startswith(prefix) for prefix in prefixes) or name.endswith("/"):
                continue
            output = target / safe_name(Path(name).name)
            with package.open(name) as source, output.open("wb") as destination:
                shutil.copyfileobj(source, destination)
            metadata = image_metadata(output)
            metadata["package_entry"] = name
            found.append(metadata)
    return found


def read_docx(path: Path) -> dict:
    document = Document(path)
    blocks = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append({"kind": "paragraph", "style": paragraph.style.name, "text": text})
    for table_index, table in enumerate(document.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        blocks.append({"kind": "table", "index": table_index, "rows": rows})
    return {
        "kind": "docx",
        "blocks": blocks,
        "media": extract_package_media(path, ("word/media/",)),
    }


def read_pptx(path: Path) -> dict:
    slides = []
    notes = []
    with zipfile.ZipFile(path) as package:
        slide_names = sorted(
            (name for name in package.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda name: int(re.search(r"(\d+)", Path(name).stem).group(1)),
        )
        for name in slide_names:
            texts = xml_text(
                package.read(name),
                "http://schemas.openxmlformats.org/drawingml/2006/main",
                "t",
            )
            slides.append({"entry": name, "text": texts})
        note_names = sorted(
            (name for name in package.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)),
            key=lambda name: int(re.search(r"(\d+)", Path(name).stem).group(1)),
        )
        for name in note_names:
            texts = xml_text(
                package.read(name),
                "http://schemas.openxmlformats.org/drawingml/2006/main",
                "t",
            )
            notes.append({"entry": name, "text": texts})
    return {
        "kind": "pptx",
        "slides": slides,
        "notes": notes,
        "media": extract_package_media(path, ("ppt/media/",)),
    }


def read_odt(path: Path) -> dict:
    with zipfile.ZipFile(path) as package:
        texts = []
        if "content.xml" in package.namelist():
            root = ET.fromstring(package.read("content.xml"))
            for node in root.iter():
                if node.text and node.text.strip():
                    texts.append(node.text.strip())
    return {
        "kind": "odt",
        "text": texts,
        "media": extract_package_media(path, ("Pictures/",)),
    }


def read_xlsx(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            if any(value is not None for value in values):
                rows.append(values)
        sheets.append({"title": sheet.title, "rows": rows})
    workbook.close()
    return {
        "kind": "xlsx",
        "sheets": sheets,
        "media": extract_package_media(path, ("xl/media/",)),
    }


def read_notebook(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    cells = []
    for index, cell in enumerate(data.get("cells", []), 1):
        source = "".join(cell.get("source", []))
        outputs = []
        for output in cell.get("outputs", []):
            text = output.get("text")
            if text:
                outputs.append("".join(text) if isinstance(text, list) else str(text))
            data_block = output.get("data", {})
            for mime in ("text/plain", "text/markdown"):
                if mime in data_block:
                    value = data_block[mime]
                    outputs.append("".join(value) if isinstance(value, list) else str(value))
        cells.append({"index": index, "type": cell.get("cell_type"), "source": source, "outputs": outputs})
    return {"kind": "ipynb", "cells": cells}


def read_text(path: Path) -> dict:
    return {"kind": "text", "text": path.read_text(encoding="utf-8", errors="replace")}


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    MEDIA.mkdir(parents=True, exist_ok=True)
    records = []
    for path in sorted((item for item in SOURCE.rglob("*") if item.is_file()), key=lambda item: str(item).lower()):
        relative = str(path.relative_to(SOURCE))
        record = {
            "path": relative,
            "extension": path.suffix.lower(),
            "bytes": path.stat().st_size,
            "modified": path.stat().st_mtime,
        }
        try:
            extension = path.suffix.lower()
            if extension == ".docx":
                record["content"] = read_docx(path)
            elif extension == ".pptx":
                record["content"] = read_pptx(path)
            elif extension == ".odt":
                record["content"] = read_odt(path)
            elif extension == ".xlsx":
                record["content"] = read_xlsx(path)
            elif extension == ".ipynb":
                record["content"] = read_notebook(path)
            elif extension in {".txt", ".md"}:
                record["content"] = read_text(path)
            elif extension in {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tif", ".tiff"}:
                record["content"] = {"kind": "image", "metadata": image_metadata(path)}
            else:
                record["content"] = {"kind": "binary"}
        except Exception as exc:
            record["error"] = f"{type(exc).__name__}: {exc}"
        records.append(record)

    (WORK / "archive-manifest.json").write_text(
        json.dumps(records, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )

    with (WORK / "archive-text.txt").open("w", encoding="utf-8") as output:
        for record in records:
            output.write(f"\n===== {record['path']} =====\n")
            content = record.get("content", {})
            kind = content.get("kind")
            if kind == "docx":
                for block in content.get("blocks", []):
                    if block["kind"] == "paragraph":
                        output.write(f"[{block['style']}] {block['text']}\n")
                    else:
                        output.write(f"[TABLE {block['index']}]\n")
                        for row in block["rows"]:
                            output.write(" | ".join(row) + "\n")
            elif kind == "pptx":
                for index, slide in enumerate(content.get("slides", []), 1):
                    output.write(f"[SLIDE {index}]\n")
                    output.write("\n".join(slide["text"]) + "\n")
                for index, note in enumerate(content.get("notes", []), 1):
                    if note["text"]:
                        output.write(f"[NOTES {index}]\n")
                        output.write("\n".join(note["text"]) + "\n")
            elif kind == "xlsx":
                for sheet in content.get("sheets", []):
                    output.write(f"[SHEET {sheet['title']}]\n")
                    for row in sheet["rows"]:
                        output.write(" | ".join("" if value is None else str(value) for value in row) + "\n")
            elif kind == "ipynb":
                for cell in content.get("cells", []):
                    output.write(f"[CELL {cell['index']} {cell['type']}]\n{cell['source']}\n")
                    for value in cell["outputs"]:
                        output.write(f"[OUTPUT]\n{value}\n")
            elif kind == "odt":
                output.write("\n".join(content.get("text", [])) + "\n")
            elif kind == "text":
                output.write(content.get("text", "") + "\n")

    print(json.dumps({"files": len(records), "errors": [r["path"] for r in records if "error" in r]}, indent=2))


if __name__ == "__main__":
    main()
